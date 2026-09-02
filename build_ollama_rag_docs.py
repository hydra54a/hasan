from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "deliverables"
OUT.mkdir(exist_ok=True)

BLUE = "2E74B5"
DARK = "1F4D78"
INK = "172B4D"
MUTED = "59636E"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
CODE_BG = "F2F4F7"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "7A5A00"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_header(row):
    set_repeat_table_header(row)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def set_run(run, font="Calibri", size=11, color="000000", bold=False, italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    return run


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), "18")
    r_pr.extend([c, u, sz, sz_cs])
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.extend([r_pr, text_node])
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def configure_document(doc, short_title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Title", 28, INK, 0, 8),
        ("Subtitle", 14, MUTED, 0, 18),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(7.6)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.0
    code_style._element.get_or_add_pPr().append(OxmlElement("w:shd"))
    code_style._element.pPr[-1].set(qn("w:fill"), CODE_BG)

    note_style = styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
    note_style.font.name = "Calibri"
    note_style.font.size = Pt(10.5)
    note_style.font.color.rgb = rgb(INK)
    note_style.paragraph_format.left_indent = Inches(0.18)
    note_style.paragraph_format.right_indent = Inches(0.18)
    note_style.paragraph_format.space_before = Pt(6)
    note_style.paragraph_format.space_after = Pt(8)
    note_style.paragraph_format.line_spacing = 1.15
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE)
    note_style._element.get_or_add_pPr().append(shd)

    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.runs[0], size=9, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    doc.core_properties.title = short_title
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.subject = "Ollama-based retrieval-augmented generation"


def add_cover(doc, kicker, title, subtitle, edition="Prepared 2 September 2026"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(86)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(kicker.upper()), size=10, color=BLUE, bold=True)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(subtitle)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    set_run(p.add_run(edition), size=10, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("LOCAL-FIRST  •  SOURCE-GROUNDED  •  DEPLOYABLE"), size=9, color=DARK, bold=True)
    doc.add_page_break()
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_run(run, font="Consolas", size=7.6, color="202124")
    return p


def add_note(doc, label, text, color=INK):
    p = doc.add_paragraph(style="Note")
    set_run(p.add_run(label + "  "), size=10.5, color=color, bold=True)
    set_run(p.add_run(text), size=10.5, color=INK)
    return p


def new_bullet_sequence(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.extend([tabs, ind])
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    existing = [int(num.get(qn("w:numId"))) for num in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(doc, text):
    num_id = getattr(doc, "_rag_bullet_num_id", None)
    if num_id is None:
        num_id = new_bullet_sequence(doc)
        doc._rag_bullet_num_id = num_id
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(text)
    return p


def new_number_sequence(doc):
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if int(num.get(qn("w:numId"))) == int(style_num_id):
            abstract_id = int(num.find(qn("w:abstractNumId")).get(qn("w:val")))
            break
    if abstract_id is None:
        abstract_id = 0
    existing = [int(num.get(qn("w:numId"))) for num in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def add_number(doc, text, num_id=None):
    p = doc.add_paragraph(style="List Number")
    if num_id is not None:
        p_pr = p._p.get_or_add_pPr()
        num_pr = p_pr.get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num_id
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), size=9.5, color=INK, bold=True)
    for values in rows:
        row = table.add_row()
        for i, text in enumerate(values):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(text), size=9.2, color="202124")
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_sources(doc, sources, page_break=False):
    heading = doc.add_heading("Sources", level=1)
    heading.paragraph_format.page_break_before = page_break
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run("Official references used for this guide. Accessed 2 September 2026."), size=10, color=MUTED)
    for label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(3)
        add_hyperlink(p, label, url)


SOURCES = [
    ("Ollama: Embeddings capability", "https://docs.ollama.com/capabilities/embeddings"),
    ("Ollama API: Generate embeddings", "https://docs.ollama.com/api/embed"),
    ("Ollama API: Generate a chat message", "https://docs.ollama.com/api/chat"),
    ("Ollama: Context length", "https://docs.ollama.com/context-length"),
    ("Ollama: Linux installation", "https://docs.ollama.com/linux"),
    ("Ollama model library: Qwen3 8B", "https://ollama.com/library/qwen3:8b"),
    ("Ollama model library: EmbeddingGemma", "https://ollama.com/library/embeddinggemma"),
    ("Google NotebookLM Help: Add or discover sources", "https://support.google.com/notebooklm/answer/16215270"),
]


def build_blueprint():
    doc = Document()
    configure_document(doc, "Ollama RAG Implementation Blueprint")
    add_cover(
        doc,
        "Technical implementation guide",
        "Ollama RAG Implementation Blueprint",
        "A private document-question-answering system using Qwen3, EmbeddingGemma, FastAPI, and a vector store",
    )

    doc.add_heading("Recommendation at a glance", level=1)
    add_note(doc, "Recommended baseline", "Use qwen3:8b to generate answers and embeddinggemma to embed both document chunks and user queries. Start with 800–1,000-character chunks, 120–150 characters of overlap, cosine similarity, and top-k = 5. Require every answer to cite retrieved source chunks.")
    p = doc.add_paragraph()
    p.add_run("Why this pair works. ").bold = True
    p.add_run("Qwen3 8B is a practical general-purpose local model (about 5.2 GB in Ollama's Q4_K_M package), while EmbeddingGemma is a compact retrieval model. Ollama explicitly recommends EmbeddingGemma, qwen3-embedding, and all-minilm for embeddings and documents that the same embedding model must be used for indexing and querying.")

    doc.add_heading("Choose the model for your hardware", level=2)
    add_table(
        doc,
        ["Host profile", "Answer model", "Embedding model", "Practical guidance"],
        [
            ("8 GB RAM, CPU-first", "qwen3:4b", "embeddinggemma", "Lowest-cost useful starting point; expect slower answers and shorter context."),
            ("12–16 GB RAM or 8–12 GB VRAM", "qwen3:8b", "embeddinggemma", "Recommended baseline for internal document Q&A."),
            ("24 GB+ RAM/VRAM", "qwen3:14b", "embeddinggemma or qwen3-embedding", "Higher answer quality; benchmark latency before rollout."),
            ("24 GB+ VRAM and quality priority", "qwen3.8:27b", "qwen3-embedding", "Current high-capability option; its Ollama package is about 18 GB, before runtime/context overhead."),
        ],
        [1700, 1550, 1900, 4210],
    )
    add_note(doc, "Sizing rule", "Model package size is not total runtime memory. Context cache and parallel requests add memory. Begin with one parallel generation, verify with ollama ps, then load-test on the actual host.", color=GOLD)

    doc.add_heading("What RAG does", level=1)
    p = doc.add_paragraph()
    p.add_run("Retrieval-augmented generation (RAG) ").bold = True
    p.add_run("does not retrain the language model on your documents. It retrieves relevant passages at question time and places those passages into the model prompt. This makes updates easier and makes source citation possible.")
    flow_numbers = new_number_sequence(doc)
    for step in (
        "Ingest: extract text and metadata from approved files.",
        "Chunk: split text into passages that preserve headings and document identity.",
        "Embed: turn every chunk into a numeric vector with EmbeddingGemma.",
        "Store: save vectors, chunk text, source name, page/section, and checksum.",
        "Retrieve: embed the user question and rank chunks by cosine similarity.",
        "Generate: give only the top passages to Qwen3 with a strict grounding prompt.",
        "Cite and log: return source identifiers and record latency, retrieved chunks, and feedback.",
    ):
        add_number(doc, step, flow_numbers)

    doc.add_heading("Reference architecture", level=2)
    add_code(doc, "User / Web UI\n    │\n    ▼\nFastAPI  ──► authentication + rate limiting\n    │\n    ├──► query embedding ──► vector search ──► top 5 chunks\n    │                              │\n    └──────────────────────────────┴──► grounded prompt ──► Qwen3:8b\n                                                     │\n                                                     └──► answer + citations")

    doc.add_heading("1. Install and verify Ollama", level=1)
    p = doc.add_paragraph()
    p.add_run("Linux installation follows the official Ollama service guidance. Review the installer in environments where piping a remote script to a shell is prohibited.")
    add_code(doc, "curl -fsSL https://ollama.com/install.sh | sh\nollama -v\nsudo systemctl status ollama\ncurl http://127.0.0.1:11434/api/tags")
    add_code(doc, "ollama pull qwen3:8b\nollama pull embeddinggemma\nollama list")
    add_note(doc, "Keep it private", "Ollama's local API does not require authentication. Bind the application to localhost or a protected internal interface. Put authentication, TLS, network restrictions, and rate limiting in front of any shared deployment.", color=RED)

    doc.add_heading("2. Create the Python service", level=1)
    add_code(doc, "python3 -m venv .venv\n. .venv/bin/activate\npip install fastapi uvicorn requests numpy pydantic python-multipart pypdf")
    p = doc.add_paragraph()
    p.add_run("Suggested project layout").bold = True
    add_code(doc, "rag-service/\n├── app/main.py          # upload and query endpoints\n├── app/extractors.py    # TXT, MD, CSV, JSON, text-PDF\n├── app/chunking.py      # heading-aware overlap chunking\n├── app/ollama.py        # /api/embed and /api/chat clients\n├── app/store.py         # vector and metadata persistence\n├── data/rag.db\n├── tests/\n└── .env.example")

    doc.add_heading("3. Implement the Ollama client", level=1)
    add_code(doc, '''import requests

OLLAMA = "http://127.0.0.1:11434"

def embed(texts: list[str]) -> list[list[float]]:
    response = requests.post(
        f"{OLLAMA}/api/embed",
        json={"model": "embeddinggemma", "input": texts},
        timeout=120,
    )
    response.raise_for_status()
    return response.json()["embeddings"]

def answer(question: str, context: str) -> dict:
    system = (
        "Answer only from CONTEXT. If the answer is absent, say: "
        "'I could not find that in the indexed sources.' Cite facts as [S1]."
    )
    prompt = f"CONTEXT:\\n{context}\\n\\nQUESTION:\\n{question}"
    payload = {
        "model": "qwen3:8b",
        "stream": False,
        "think": False,
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": prompt},
        ],
        "options": {
            "temperature": 0.1,
            "num_ctx": 8192,
        },
    }
    response = requests.post(
        f"{OLLAMA}/api/chat",
        json=payload,
        timeout=300,
    )
    response.raise_for_status()
    return response.json()
''')
    add_note(doc, "API fact", "Ollama's /api/embed endpoint accepts either one string or an array of strings. Its embedding capability guide says returned vectors are L2-normalized, so cosine similarity is a natural ranking method.")

    doc.add_heading("4. Chunk and index documents", level=1)
    add_bullet(doc, "Keep source_id, filename, title, heading path, page number when available, chunk index, checksum, and updated timestamp with every chunk.")
    add_bullet(doc, "Prefer paragraph and heading boundaries; do not cut in the middle of a sentence when a nearby boundary exists.")
    add_bullet(doc, "Start with 800–1,000 characters and 120–150 characters of overlap. Tune using evaluation questions from real users.")
    add_bullet(doc, "Batch embedding requests. Re-embed only changed documents, identified by checksum.")
    add_bullet(doc, "Use the exact same embedding model for indexing and querying. If you change models, rebuild the entire index.")
    add_code(doc, '''import numpy as np

def cosine_scores(query_vector, document_matrix):
    q = np.asarray(query_vector, dtype=np.float32)
    d = np.asarray(document_matrix, dtype=np.float32)
    q /= np.linalg.norm(q) + 1e-12
    d /= np.linalg.norm(d, axis=1, keepdims=True) + 1e-12
    return d @ q
''')

    doc.add_heading("5. Retrieval and grounded answering", level=1)
    retrieval_numbers = new_number_sequence(doc)
    add_number(doc, "Normalize the user question and reject empty or oversized input.", retrieval_numbers)
    add_number(doc, "Embed the question with EmbeddingGemma.", retrieval_numbers)
    add_number(doc, "Search vectors and take the top 8–12 candidates.", retrieval_numbers)
    add_number(doc, "Optionally apply metadata filters and a reranker; keep the best 4–6 passages.", retrieval_numbers)
    add_number(doc, "Label each passage [S1], [S2], and so on, and include the source metadata in the prompt.", retrieval_numbers)
    add_number(doc, "Require citations next to supported claims. Return an explicit not-found answer when evidence is missing.", retrieval_numbers)
    add_number(doc, "Validate that every cited identifier was actually supplied to the model before returning the answer.", retrieval_numbers)

    doc.add_heading("Grounding prompt", level=2)
    add_code(doc, "You are a document assistant. Use only the supplied CONTEXT.\n\nRules:\n1. Do not use outside knowledge.\n2. Cite each factual claim using [S1], [S2], etc.\n3. If sources conflict, name the conflict and cite both.\n4. If the answer is not in context, say exactly:\n   I could not find that in the indexed sources.\n5. Do not invent source names, page numbers, policies, dates, or quotations.")

    doc.add_heading("6. API contract", level=1)
    add_table(
        doc,
        ["Endpoint", "Purpose", "Minimum response"],
        [
            ("POST /documents", "Upload and index a supported document", "document_id, status, chunks, checksum"),
            ("GET /documents", "List indexed documents", "id, filename, updated_at, chunk_count"),
            ("DELETE /documents/{id}", "Remove document and its chunks", "deleted=true"),
            ("POST /query", "Retrieve and generate an answer", "answer, sources, timings, model"),
            ("GET /health", "Check application and Ollama readiness", "database, ollama, models"),
        ],
        [1900, 3400, 4060],
    )
    add_code(doc, '''{
  "question": "What is the retention period?",
  "answer": "Records are retained for five years [S1].",
  "sources": [
    {"id": "S1", "document": "records-policy.pdf", "page": 7,
     "excerpt": "...retained for five years...", "score": 0.83}
  ],
  "model": "qwen3:8b",
  "embedding_model": "embeddinggemma"
}''')

    doc.add_heading("7. Deployment baseline", level=1)
    add_bullet(doc, "Run the API as a dedicated non-login service account and store application data outside the code directory.")
    add_bullet(doc, "Bind FastAPI and Ollama to localhost when a reverse proxy is on the same host.")
    add_bullet(doc, "Use TLS and authenticated access; restrict CORS to the actual frontend origin.")
    add_bullet(doc, "Enforce file type, byte-size, page-count, and extracted-text limits before indexing.")
    add_bullet(doc, "Scan uploads, prevent path traversal, and never execute embedded content or macros.")
    add_bullet(doc, "Back up document metadata and the source files. A vector index can be rebuilt; authorization and audit records are more critical.")
    add_bullet(doc, "Separate collections or enforce row-level access when users must not retrieve each other's documents.")

    doc.add_heading("8. Acceptance tests", level=1)
    for test in (
        "Known-answer: a question with one clear passage returns the correct fact and source.",
        "Not-found: a question outside the corpus returns the explicit not-found response.",
        "Conflict: two documents with different values are both cited and the conflict is stated.",
        "Deletion: deleting a document removes its chunks from future retrieval.",
        "Update: changing a document checksum replaces old chunks instead of duplicating them.",
        "Access control: a user cannot retrieve chunks from an unauthorized collection.",
        "Injection resistance: instructions embedded inside a source are treated as content, not system commands.",
        "Citation integrity: every returned source ID exists among retrieved chunks.",
        "Performance: record retrieval time, time to first token, total time, and generated tokens per second.",
    ):
        add_bullet(doc, test)

    doc.add_heading("9. Operational measurements", level=1)
    p = doc.add_paragraph()
    p.add_run("From Ollama's chat response, record ").bold = True
    p.add_run("total_duration, load_duration, prompt_eval_count, prompt_eval_duration, eval_count, and eval_duration. Approximate generation throughput as:")
    add_code(doc, "tokens_per_second = eval_count / (eval_duration / 1_000_000_000)")
    p = doc.add_paragraph()
    p.add_run("Concurrency warning. ").bold = True
    p.add_run("Parallel requests increase context memory. Start with OLLAMA_NUM_PARALLEL=1, then raise it only after measuring memory, queuing, and latency on the real host.")

    doc.add_heading("10. Go-live checklist", level=1)
    for item in (
        "Ollama, qwen3:8b, and embeddinggemma are installed and visible in ollama list.",
        "Health checks fail closed when Ollama or the vector store is unavailable.",
        "Documents are versioned, checksummed, and access-controlled.",
        "Answers are grounded, source IDs are validated, and not-found behavior is tested.",
        "TLS, authentication, authorization, rate limiting, upload controls, backups, and audit logs are enabled.",
        "A representative evaluation set meets agreed accuracy, citation, latency, and security targets.",
        "The model and context fit fully in the intended RAM/VRAM under expected concurrency.",
    ):
        add_bullet(doc, item)

    add_sources(doc, SOURCES[:-1])
    path = OUT / "Ollama_RAG_Implementation_Blueprint.docx"
    doc.save(path)
    return path


def build_notebook_pack():
    doc = Document()
    configure_document(doc, "NotebookLM Source Pack for an Ollama RAG Project")
    add_cover(
        doc,
        "NotebookLM-ready study pack",
        "Ollama RAG Project Source Pack",
        "Upload this document to NotebookLM to plan, review, and test a private retrieval-augmented generation implementation",
    )

    doc.add_heading("How to use this source", level=1)
    p = doc.add_paragraph()
    p.add_run("Upload this DOCX to NotebookLM together with ").bold = True
    p.add_run("the implementation blueprint, your own policies/manuals, the Ollama API pages listed at the end, and a system requirements document. NotebookLM accepts DOCX, Markdown, TXT, PDF, CSV, PPTX, web URLs, and several other source types; Google states that each uploaded source may contain up to 500,000 words or 200 MB, subject to plan limits.")
    add_note(doc, "Important distinction", "NotebookLM helps you study, compare, and reason over the project sources. It does not become the production vector database or the Ollama runtime. The production RAG service still needs its own ingestion, embeddings, retrieval, access control, and evaluation.")

    doc.add_heading("Project decision record", level=1)
    add_table(
        doc,
        ["Decision", "Selected baseline", "Reason"],
        [
            ("Answer model", "qwen3:8b", "Practical local model for 12–16 GB-class hosts; multilingual and instruction-following capable."),
            ("Embedding model", "embeddinggemma", "Compact Ollama-recommended embedding model for retrieval and semantic similarity."),
            ("API", "FastAPI", "Simple typed endpoints for upload, indexing, querying, health, and administration."),
            ("Vector store", "SQLite + vectors for prototype; pgvector or Qdrant for scale", "Start simple, then add stronger filtering, concurrency, and operations when needed."),
            ("Chunking", "800–1,000 chars; 120–150 overlap", "A testable starting point, not a universal optimum."),
            ("Retrieval", "cosine similarity; top-k 5 after candidate filtering", "Matches normalized text embeddings and keeps prompt context focused."),
            ("Grounding", "context-only answers with [S#] citations", "Makes unsupported answers visible and testable."),
        ],
        [2000, 2700, 4660],
    )

    doc.add_heading("RAG vocabulary", level=1)
    terms = [
        ("Corpus", "The complete approved set of source documents available to the system."),
        ("Chunk", "A retrievable passage derived from a source, stored with metadata and a stable identifier."),
        ("Embedding", "A numeric vector representing text meaning for similarity search."),
        ("Vector store", "A database or index that stores embeddings and lets the application search for similar vectors."),
        ("Retriever", "The component that selects passages relevant to the user's question."),
        ("Top-k", "The number of highest-ranked passages retained at a retrieval stage."),
        ("Reranker", "An optional second-stage model that reorders candidate passages for relevance."),
        ("Grounding", "Constraining the generated answer to retrieved evidence."),
        ("Hallucination", "A confident statement that is unsupported, incorrect, or invented."),
        ("Citation integrity", "The requirement that every cited source was retrieved and actually supports the adjacent claim."),
    ]
    for term, definition in terms:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run(p.add_run(term + ": "), bold=True, color=DARK)
        p.add_run(definition)

    doc.add_heading("Documents to add to your NotebookLM notebook", level=1)
    for item in (
        "This source pack and the Ollama RAG Implementation Blueprint.",
        "The official Ollama Embeddings, /api/embed, /api/chat, Context Length, Linux, and selected model library pages.",
        "Your business requirements: intended users, document types, expected question volume, answer latency, and availability target.",
        "Your security and privacy requirements: data classification, access rules, retention, deletion, audit, backup, and incident response.",
        "A document inventory with owner, sensitivity, format, update frequency, and authoritative version.",
        "A representative evaluation set with questions, expected facts, required source passages, and not-found cases.",
        "Hardware specifications and measured Ollama performance from the intended deployment host.",
        "A change log recording model, embedding model, chunking, retrieval, prompt, and index versions.",
    ):
        add_bullet(doc, item)

    doc.add_heading("Source preparation rules", level=1)
    add_bullet(doc, "Use authoritative, current, and legally permitted sources. Record owner and effective date.")
    add_bullet(doc, "Prefer text-searchable files. OCR image-only PDFs before ingestion and spot-check the extracted text.")
    add_bullet(doc, "Keep headings, section numbers, page numbers, table labels, and document titles; they improve metadata and citations.")
    add_bullet(doc, "Remove duplicate and obsolete copies or label superseded versions clearly.")
    add_bullet(doc, "Do not upload passwords, API keys, private tokens, or unrelated personal data.")
    add_bullet(doc, "For production RAG, enforce authorization at retrieval time; source filtering only in the user interface is insufficient.")

    doc.add_heading("Questions to ask NotebookLM", level=1)
    prompts = [
        "Using only the selected sources, summarize the proposed RAG architecture and list every component that must be implemented.",
        "Compare qwen3:4b, qwen3:8b, qwen3:14b, and qwen3.8:27b for our stated hardware. Separate documented facts from assumptions.",
        "Create a requirements traceability matrix linking each security requirement to an implementation control and an acceptance test.",
        "Find contradictions among the source documents about model choice, chunk size, context length, or deployment topology. Cite both sides.",
        "Create 30 evaluation questions: 15 direct facts, 5 multi-section synthesis questions, 5 conflicting-source questions, and 5 not-found questions.",
        "For every evaluation question, identify the expected source document and exact supporting passage. Do not invent missing answers.",
        "Review the API contract and identify missing error cases, authorization checks, observability fields, and deletion behavior.",
        "Create a phased implementation plan: prototype, internal pilot, security hardening, load test, and production release.",
        "Draft a failure-mode review for document parsing, embedding, retrieval, generation, citation, model availability, and database corruption.",
        "Produce a go-live checklist grounded only in the uploaded requirements and implementation blueprint.",
    ]
    prompt_numbers = new_number_sequence(doc)
    for prompt in prompts:
        add_number(doc, prompt, prompt_numbers)

    doc.add_heading("Evaluation worksheet", level=1)
    add_table(
        doc,
        ["Field", "What to record"],
        [
            ("Question ID", "Stable identifier, user role, collection, and question text"),
            ("Expected evidence", "Authoritative source, page/section, and supporting passage"),
            ("Expected behavior", "Answer facts or explicit not-found response"),
            ("Retrieval result", "Top chunk IDs, scores, filters, and whether required evidence was retrieved"),
            ("Answer result", "Correctness, completeness, unsupported claims, and conflict handling"),
            ("Citation result", "Citation presence, validity, and whether each passage supports its claim"),
            ("Performance", "Retrieval time, first-token time, total time, and tokens/second"),
            ("Outcome", "Pass/fail, defect category, owner, and corrective action"),
        ],
        [2300, 7060],
    )

    doc.add_heading("Minimum acceptance criteria", level=1)
    for criterion in (
        "The system retrieves the required supporting passage for agreed known-answer questions.",
        "Answers contain no unsupported factual claims in the acceptance set.",
        "Every citation maps to a retrieved chunk and supports the adjacent claim.",
        "Not-found questions produce an explicit not-found response rather than a plausible guess.",
        "Unauthorized users cannot retrieve protected chunks, including through indirect or adversarial prompts.",
        "Document deletion and replacement remove obsolete chunks from retrieval.",
        "Latency and memory remain within agreed limits under measured concurrent load.",
        "Logs support incident review without exposing unnecessary document content or secrets.",
    ):
        add_bullet(doc, criterion)

    doc.add_heading("Information still needed before final sizing", level=1)
    add_bullet(doc, "Operating system, CPU model/core count, total RAM, GPU model, and VRAM.")
    add_bullet(doc, "Number of documents, total extracted text size, languages, and scanned-PDF percentage.")
    add_bullet(doc, "Expected simultaneous users, requests per minute, typical answer length, and latency target.")
    add_bullet(doc, "Whether the corpus contains confidential or regulated information.")
    add_bullet(doc, "Required authentication provider, user roles, and document-level access rules.")
    add_bullet(doc, "Availability, backup, recovery-time, retention, deletion, and audit requirements.")

    add_sources(doc, SOURCES, page_break=True)
    path = OUT / "NotebookLM_Ollama_RAG_Source_Pack.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for created in (build_blueprint(), build_notebook_pack()):
        print(created)
